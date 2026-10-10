from __future__ import annotations

import asyncio
import importlib
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

models = importlib.import_module("Turning-Good-Agent.web.attachments.models")
documents = importlib.import_module("Turning-Good-Agent.web.attachments.documents")

AttachmentMetadata = models.AttachmentMetadata
AttachmentOwner = models.AttachmentOwner
ResolvedAttachment = models.ResolvedAttachment
read_document = documents.read_document
AttachmentReadError = documents.AttachmentReadError


def resolved(tmp_path: Path, filename: str, content: bytes) -> ResolvedAttachment:
    path = tmp_path / filename
    path.write_bytes(content)
    metadata = AttachmentMetadata(
        "att_test", "document", filename, "text/plain", len(content)
    )
    return ResolvedAttachment(metadata, path, "digest", AttachmentOwner("p", "web", "c"))


def test_text_reader_returns_one_based_line_numbers(tmp_path):
    result = read_document(resolved(tmp_path, "note.txt", b"a\nb\nc\n"), offset=2, limit=2)
    assert result.content == "2|b\n3|c"
    assert result.range_label == "lines 2-3"
    assert result.has_more is False
    assert result.truncated is False


def test_text_reader_reports_more_lines(tmp_path):
    result = read_document(resolved(tmp_path, "note.txt", b"a\nb\nc\n"), offset=1, limit=2)
    assert result.has_more is True


def test_text_reader_bounds_very_long_lines(tmp_path):
    result = read_document(resolved(tmp_path, "note.txt", ("x" * 140_000).encode()))
    assert result.truncated is True
    assert result.has_more is True
    assert len(result.content) < 129_000


def test_pdf_reader_accepts_a_real_pdf_and_reports_page_range(tmp_path):
    path = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)
    attachment = ResolvedAttachment(
        AttachmentMetadata("att_test", "document", path.name, "application/pdf", path.stat().st_size),
        path,
        "digest",
        AttachmentOwner("p", "web", "c"),
    )

    result = read_document(attachment, pages="2-2")

    assert result.range_label == "pages 2-2"
    assert result.total_units == 2


@pytest.mark.parametrize("suffix, expected", [(".docx", "DOCX evidence"), (".xlsx", "XLSX evidence"), (".pptx", "PPTX evidence")])
def test_office_readers_extract_real_document_text(tmp_path, suffix, expected):
    path = tmp_path / f"sample{suffix}"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph(expected)
        document.save(path)
    elif suffix == ".xlsx":
        workbook = Workbook()
        workbook.active["A1"] = expected
        workbook.save(path)
        workbook.close()
    else:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = expected
        presentation.save(path)
    attachment = ResolvedAttachment(
        AttachmentMetadata("att_test", "document", path.name, "application/octet-stream", path.stat().st_size),
        path,
        "digest",
        AttachmentOwner("p", "web", "c"),
    )

    result = read_document(attachment)

    assert expected in result.content
    assert result.truncated is False


def test_unsupported_image_is_rejected(tmp_path):
    attachment = resolved(tmp_path, "photo.png", b"png")
    attachment = ResolvedAttachment(
        AttachmentMetadata("att_test", "image", "photo.png", "image/png", 3),
        attachment.path,
        attachment.digest,
        attachment.owner,
    )
    with pytest.raises(AttachmentReadError, match="图片不能通过文档读取"):
        read_document(attachment)


def test_invalid_pdf_page_range_is_structured(tmp_path):
    attachment = resolved(tmp_path, "sample.pdf", b"%PDF-not-valid")
    attachment = ResolvedAttachment(
        AttachmentMetadata("att_test", "document", "sample.pdf", "application/pdf", 14),
        attachment.path,
        attachment.digest,
        attachment.owner,
    )
    with pytest.raises(AttachmentReadError):
        read_document(attachment, pages="bad")
